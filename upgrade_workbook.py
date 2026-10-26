#!/usr/bin/env python3
"""
upgrade_workbook.py - bring a topology workbook up to the current format (and refresh dropdowns)

  python upgrade_workbook.py topology.xlsx [more.xlsx ...]

What it does (a backup <name>.bak.xlsx is written first):
  * Links: drops the old 'type' / 'media' columns and fills 'cable' (fiber | eth | dac) from them;
    puts the columns in the standard order; keeps every row and any extra columns you added
  * Devices: adds a 'kind' column (switch | ixia | server | mgmt) if missing, default switch
  * replaces the old LinkTypes sheet with Roles (colour per connection kind) and Cables
  * refreshes the HowTo sheet and all dropdowns
Safe to run more than once.
"""
import shutil
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("This script needs openpyxl:  pip install openpyxl")

LINK_COLS = ["device_a", "port_a", "patch_a", "device_b", "port_b", "patch_b", "cable", "speed",
             "status", "date_changed", "notes", "loop_side"]
LINK_WIDTHS = [12, 9, 8, 12, 9, 8, 8, 8, 10, 13, 34, 10]
DEVICE_COLS = ["name", "kind", "model", "role", "emp_ip", "tier", "notes"]
ROLES = [
    ["dut-dut", "#1f3a5f", "Switch <-> switch"],
    ["traffic", "#2a6fbd", "Switch -> server NIC (arrow points at the server)"],
    ["ixia", "#2e9e5b", "Switch <-> Ixia port"],
    ["mgmt", "#8b95a5", "Anything <-> management network"],
    ["loopback", "#7c3aed", "Port to port on the same device"],
    ["other", "#6b7280", "Any other pair (e.g. server <-> Ixia)"],
]
CABLES = [
    ["eth", 2, "solid", "no", "Copper Ethernet (RJ45)"],
    ["fiber", 3, "solid", "yes", "Optical fiber"],
    ["dac", 4, "double", "yes", "Direct-attach copper (twinax)"],
]
SPEEDS = "1G,2.5G,10G,25G,40G,50G,100G,200G,400G"
HOWTO = [
    "How to use",
    "Devices: one row per box. kind = switch | ixia | server | mgmt. tier = row in the diagram (1 = top); "
    "left-to-right order = sheet order.",
    "Links: one row per cable. Pick device names from the dropdown. Leave a port blank if it is not known "
    "(the line is then drawn dotted).",
    "cable = fiber | eth | dac.  speed = 10G, 40G, 100G ... (shown on the line for fiber and DAC).",
    "You never enter the connection kind: switch-switch = dut-dut, switch-server = traffic, switch-ixia = ixia, "
    "anything-mgmt = mgmt, same device both ends = loopback.",
    "Colours: Roles sheet (one colour per connection kind). Line thickness/style: Cables sheet.",
    "patch_a / patch_b: patch-panel port (P#) at each end. Neighbouring links with the same P# share one bracket.",
    "Add a connection: new row, status = active. Remove: status = removed (keeps history) or delete the row.",
    "status = planned is hidden unless you run: python topo.py <workbook> --show-planned",
    "loop_side (optional, loopbacks only) = top | bottom | left | right.",
    "Run: python topo.py <workbook> --title \"My setup\"  ->  .svg + .drawio next to the workbook",
]
HDR_FONT, HDR_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
ALIASES = {"ethernet": "eth", "copper": "eth", "rj45": "eth", "fibre": "fiber", "optic": "fiber",
           "optical": "fiber", "twinax": "dac"}


def val(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    hdr = [val(h).lower().replace(" ", "_") for h in rows[0]]
    data = [dict(zip(hdr, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]
    return hdr, data


def write(wb, name, cols, rows, widths, index=None):
    if name in wb.sheetnames:
        index = wb.sheetnames.index(name) if index is None else index
        del wb[name]
    ws = wb.create_sheet(name, index)
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = HDR_FONT, HDR_FILL
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"
    return ws


def dropdown(ws, formula, col_letters, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showErrorMessage = strict
    ws.add_data_validation(dv)
    for c in col_letters:
        dv.add(f"{c}2:{c}1000")


def cable_from(r):
    c = val(r.get("cable")).lower()
    if not c:
        c = val(r.get("media")).lower()
    if not c:
        t = val(r.get("type")).lower()
        c = "fiber" if "fiber" in t or "fibre" in t else "dac" if "dac" in t else ""
    return ALIASES.get(c, c)


def upgrade(path):
    path = Path(path)
    shutil.copy2(path, path.with_name(path.stem + ".bak.xlsx"))
    wb = openpyxl.load_workbook(path)

    # Devices
    hdr, devs = read(wb["Devices"])
    extra = [h for h in hdr if h and h not in DEVICE_COLS]
    cols = DEVICE_COLS + extra
    rows = [[val(d.get(c)) if c != "kind" else (val(d.get("kind")).lower() or "switch") for c in cols] for d in devs]
    for r in rows:
        ti = cols.index("tier")
        r[ti] = int(r[ti]) if r[ti].isdigit() else r[ti]
    ws = write(wb, "Devices", cols, rows, [14, 9, 14, 10, 15, 6, 40] + [14] * len(extra))
    dropdown(ws, '"switch,ixia,server,mgmt"', ["B"])

    # Links
    hdr, links = read(wb["Links"])
    extra = [h for h in hdr if h and h not in LINK_COLS and h not in ("type", "media")]
    cols = LINK_COLS + extra
    rows = []
    for r in links:
        r = dict(r)
        r["cable"] = cable_from(r)
        if not val(r.get("status")):
            r["status"] = "active"
        rows.append([val(r.get(c)) for c in cols])
    ws = write(wb, "Links", cols, rows, LINK_WIDTHS + [14] * len(extra))
    dropdown(ws, "=Devices!$A$2:$A$300", ["A", "D"])
    dropdown(ws, "=Cables!$A$2:$A$50", ["G"])
    dropdown(ws, f'"{SPEEDS}"', ["H"], strict=False)       # other speeds can still be typed
    dropdown(ws, '"active,planned,removed"', ["I"])
    dropdown(ws, '"top,bottom,left,right"', ["L"])

    # Roles / Cables (keep colours / widths the user already changed)
    if "LinkTypes" in wb.sheetnames:
        del wb["LinkTypes"]
    roles = [list(r) for r in ROLES]
    if "Roles" in wb.sheetnames:
        _, old = read(wb["Roles"])
        mine = {val(r.get("role")).lower(): val(r.get("color")) for r in old}
        for r in roles:
            r[1] = mine.get(r[0]) or r[1]
    write(wb, "Roles", ["role", "color", "description"], roles, [10, 10, 52])
    cables = [list(c) for c in CABLES]
    if "Cables" in wb.sheetnames:
        _, old = read(wb["Cables"])
        known = {c[0] for c in cables}
        for o in old:
            k = val(o.get("cable")).lower()
            row = [k, o.get("width") or 2, val(o.get("line")) or "solid", val(o.get("show_speed")) or "no",
                   val(o.get("description"))]
            if k in known:
                cables[[c[0] for c in cables].index(k)] = row
            elif k:
                cables.append(row)
    ws = write(wb, "Cables", ["cable", "width", "line", "show_speed", "description"], cables, [8, 7, 9, 11, 34])
    dropdown(ws, '"solid,double,dashdot"', ["C"])
    dropdown(ws, '"yes,no"', ["D"])

    # HowTo first, Links active
    if "HowTo" in wb.sheetnames:
        del wb["HowTo"]
    ws = wb.create_sheet("HowTo", 0)
    for line in HOWTO:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 125
    order = ["HowTo", "Devices", "Links", "Roles", "Cables"]
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
    wb.active = wb.sheetnames.index("Links")
    for s in wb.worksheets:
        s.sheet_view.tabSelected = s.title == "Links"
    wb.save(path)
    print(f"{path.name}: upgraded ({len(rows)} links); backup in {path.stem}.bak.xlsx")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    for p in sys.argv[1:]:
        upgrade(p)
